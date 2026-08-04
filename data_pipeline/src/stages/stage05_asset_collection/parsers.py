from __future__ import annotations

import csv
import json
import os
import re
import subprocess
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import yaml

from src.core.io import write_json
from src.integrations.mineru import run_mineru_queue
from src.stages.stage05_asset_collection.archive import is_archive, safe_extract
from src.stages.stage05_asset_collection.clues import extract_clues

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".rst",
    ".log",
    ".out",
    ".inp",
    ".in",
    ".com",
    ".gjf",
    ".py",
    ".sh",
    ".xml",
    ".html",
    ".htm",
    ".cif",
    ".xyz",
    ".pdb",
    ".mol",
    ".sdf",
    ".gro",
    ".top",
}


def parse_asset(
    asset: dict[str, Any],
    *,
    parsed_root: Path,
    mineru_config: dict[str, Any],
    next_round: int,
    max_text_chars: int = 2_000_000,
    max_archive_files: int = 5000,
    max_archive_bytes: int = 50 * 1024**3,
) -> tuple[dict[str, Any], list[Path], list[dict[str, Any]]]:
    updated = dict(asset)
    source = Path(str(asset["original_path"]))
    output = parsed_root / str(asset["asset_id"])
    output.mkdir(parents=True, exist_ok=True)
    file_name = str(asset.get("file_name") or source.name)
    suffix = _suffix(file_name)
    parser_source = _named_source(source, output, file_name)
    children: list[Path] = []
    extracted_clues: list[dict[str, Any]] = []
    text = ""
    structured: dict[str, Any] = {
        "asset_id": asset["asset_id"],
        "file_name": file_name,
        "media_type": asset.get("media_type"),
        "size_bytes": asset.get("size_bytes"),
        "sha256": asset.get("sha256"),
    }
    try:
        if suffix == ".pdf":
            text, pdf_data = _parse_pdf(
                {**asset, "original_path": str(parser_source)}, output, mineru_config
            )
            structured.update(pdf_data)
            updated["parser"] = pdf_data.get("parser")
        elif is_archive(parser_source, file_name):
            children = safe_extract(
                parser_source,
                output / "extracted",
                max_files=max_archive_files,
                max_total_bytes=max_archive_bytes,
            )
            structured.update({"parser": "archive", "child_files": len(children)})
            text = "\n".join(str(item.relative_to(output / "extracted")) for item in children)
            updated["parser"] = "archive"
        elif suffix in {".json", ".yaml", ".yml"}:
            raw = parser_source.read_text(encoding="utf-8", errors="replace")[:max_text_chars]
            payload = json.loads(raw) if suffix == ".json" else yaml.safe_load(raw)
            structured.update({"parser": "structured_text", "value": payload})
            text = json.dumps(payload, ensure_ascii=False, indent=2)
            updated["parser"] = "structured_text"
        elif suffix in {".csv", ".tsv"}:
            text, table = _parse_table(parser_source, suffix, max_text_chars)
            structured.update(table)
            updated["parser"] = "table"
        elif suffix == ".docx":
            text = _parse_docx(parser_source, max_text_chars)
            structured.update({"parser": "docx", "characters": len(text)})
            updated["parser"] = "docx"
        elif suffix == ".xlsx":
            text, workbook = _parse_xlsx(parser_source, max_text_chars)
            structured.update(workbook)
            updated["parser"] = "xlsx"
        elif suffix in TEXT_SUFFIXES or str(asset.get("media_type", "")).startswith("text/"):
            raw = parser_source.read_text(encoding="utf-8", errors="replace")[:max_text_chars]
            is_html = suffix in {".html", ".htm"} or asset.get("media_type") == "text/html"
            if is_html:
                text, extracted_clues = _parse_html(
                    raw,
                    base_url=str(asset.get("resolved_url") or asset.get("source_url") or ""),
                    paper_id=str(asset["paper_id"]),
                    next_round=next_round,
                    source_asset_id=str(asset["asset_id"]),
                )
            else:
                text = raw
            if suffix == ".xml":
                text = re.sub(r"<[^>]+>", " ", text)
            structured.update({"parser": "text", "characters": len(text)})
            updated["parser"] = "text"
        else:
            structured.update({"parser": "binary_metadata", "note": "Original binary preserved"})
            text = f"Binary asset preserved: {file_name}\nSHA-256: {asset.get('sha256')}\n"
            updated["parser"] = "binary_metadata"
        text = text[:max_text_chars]
        markdown_path = output / "content.md"
        text_path = output / "content.txt"
        markdown_path.write_text(text, encoding="utf-8")
        text_path.write_text(text, encoding="utf-8")
        structured_path = write_json(output / "document.json", structured)
        updated["parse_status"] = "success" if text or children else "partial"
        updated["structured_path"] = str(structured_path.resolve())
        updated["readable_paths"] = [str(markdown_path.resolve()), str(text_path.resolve())]
        updated["error"] = None
    except Exception as exc:
        updated["parse_status"] = "failed"
        updated["error"] = f"{type(exc).__name__}: {exc}"
        write_json(output / "document.json", {**structured, "error": updated["error"]})
    clues = extracted_clues or extract_clues(
        text,
        paper_id=str(asset["paper_id"]),
        discovery_round=next_round,
        source_asset_id=str(asset["asset_id"]),
        max_clues=200,
    )
    return updated, children, clues


def _named_source(source: Path, output: Path, file_name: str) -> Path:
    safe_name = Path(file_name).name or source.name
    target = output / safe_name
    if target.exists() or target.is_symlink():
        return target
    try:
        os.symlink(source, target)
    except OSError:
        return source
    return target


def _parse_pdf(
    asset: dict[str, Any], output: Path, mineru_config: dict[str, Any]
) -> tuple[str, dict[str, Any]]:
    page_count = _pdf_page_count(Path(str(asset["original_path"])))
    max_pages = int(mineru_config.get("max_pages", 100))
    if page_count is not None and page_count > max_pages:
        text = _pdftotext(
            Path(str(asset["original_path"])),
            output / "pdftotext.txt",
            timeout_seconds=int(mineru_config.get("pdftotext_timeout_seconds", 600)),
        )
        return text, {
            "parser": "pdftotext_large_pdf",
            "page_count": page_count,
            "mineru_skipped": True,
            "mineru_skip_reason": f"page_count {page_count} exceeds max_pages {max_pages}",
        }
    result = run_mineru_queue(
        [
            {
                "document_id": asset["asset_id"],
                "paper_id": asset["paper_id"],
                "source_path": asset["original_path"],
                "title": asset.get("file_name"),
                "expected_pages": None,
                "deep_parse_decision": "required",
            }
        ],
        output / "mineru",
        execute=bool(mineru_config.get("execute", True)),
        command=str(mineru_config.get("command", "mineru")),
        method=str(mineru_config.get("method", "auto")),
        backend=mineru_config.get("backend"),
        timeout_seconds=int(mineru_config.get("timeout_seconds", 3600)),
        working_directory=mineru_config.get("working_directory"),
        environment=mineru_config.get("environment"),
        extra_args=mineru_config.get("extra_args"),
        reuse_existing=bool(mineru_config.get("reuse_existing", True)),
        min_markdown_chars=int(mineru_config.get("min_markdown_chars", 100)),
        stage_name="stage_05_asset_collection",
    )[0]
    markdown = Path(str(result.get("markdown_path") or ""))
    if result.get("status") in {"success", "reused"} and markdown.is_file():
        return markdown.read_text(encoding="utf-8", errors="replace"), {
            "parser": "mineru",
            "mineru": result,
        }
    fallback = Path(str(asset.get("fallback_text_path") or ""))
    if fallback.is_file():
        return fallback.read_text(encoding="utf-8", errors="replace"), {
            "parser": "grobid_fallback",
            "mineru": result,
        }
    raise RuntimeError(result.get("error") or f"MinerU status: {result.get('status')}")


def _pdf_page_count(path: Path) -> int | None:
    try:
        completed = subprocess.run(
            ["pdfinfo", str(path)],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return None
    if completed.returncode != 0:
        return None
    match = re.search(r"^Pages:\s*(\d+)\s*$", completed.stdout, re.MULTILINE)
    return int(match.group(1)) if match else None


def _pdftotext(path: Path, output: Path, *, timeout_seconds: int) -> str:
    completed = subprocess.run(
        ["pdftotext", "-layout", str(path), str(output)],
        capture_output=True,
        text=True,
        timeout=timeout_seconds,
        check=False,
    )
    if completed.returncode != 0 or not output.is_file():
        error = completed.stderr[-1000:] or f"pdftotext exited {completed.returncode}"
        raise RuntimeError(error)
    text = output.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        raise RuntimeError("pdftotext produced empty output")
    return text


def _parse_table(path: Path, suffix: str, max_chars: int) -> tuple[str, dict[str, Any]]:
    delimiter = "\t" if suffix == ".tsv" else ","
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        for index, row in enumerate(csv.reader(handle, delimiter=delimiter)):
            rows.append(row)
            if index >= 99:
                break
    text = "\n".join(delimiter.join(row) for row in rows)[:max_chars]
    return text, {
        "parser": "table",
        "preview_rows": len(rows),
        "columns": max((len(row) for row in rows), default=0),
    }


def _parse_docx(path: Path, max_chars: int) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)[:max_chars]


def _parse_xlsx(path: Path, max_chars: int) -> tuple[str, dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        count = 0
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if value is None else str(value) for value in row))
            count += 1
            if count >= 100 or sum(len(item) for item in lines) >= max_chars:
                break
        sheets.append({"name": sheet.title, "preview_rows": count, "max_column": sheet.max_column})
    workbook.close()
    return "\n".join(lines)[:max_chars], {"parser": "xlsx", "sheets": sheets}


class _LinkTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self.text: list[str] = []
        self._href: str | None = None
        self._anchor_text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() == "a":
            self._href = next((value for key, value in attrs if key == "href" and value), None)
            self._anchor_text = []

    def handle_data(self, data: str) -> None:
        self.text.append(data)
        if self._href:
            self._anchor_text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() == "a" and self._href:
            self.links.append((self._href, " ".join(self._anchor_text)))
            self._href = None
            self._anchor_text = []


def _parse_html(
    raw: str,
    *,
    base_url: str,
    paper_id: str,
    next_round: int,
    source_asset_id: str,
) -> tuple[str, list[dict[str, Any]]]:
    parser = _LinkTextParser()
    parser.feed(raw)
    text = " ".join(" ".join(parser.text).split())
    clues: list[dict[str, Any]] = []
    terms = (
        "supplement",
        "supporting information",
        "source data",
        "download data",
        "download code",
        "repository",
    )
    for href, anchor_text in parser.links:
        url = urljoin(base_url, href)
        label = " ".join(anchor_text.split())
        lower = label.casefold()
        if not any(term in lower for term in terms) and not _asset_link(url):
            continue
        clues.extend(
            extract_clues(
                f"Supporting information is available at {url}. {label}",
                paper_id=paper_id,
                discovery_round=next_round,
                source_asset_id=source_asset_id,
                max_clues=20,
            )
        )
    return text, clues


def _asset_link(url: str) -> bool:
    lower = url.casefold().split("?", 1)[0]
    return any(
        lower.endswith(suffix)
        for suffix in (
            ".zip",
            ".tar",
            ".tar.gz",
            ".tgz",
            ".pdf",
            ".csv",
            ".tsv",
            ".xlsx",
            ".json",
            ".yaml",
            ".yml",
            ".cif",
            ".xyz",
            ".pdb",
            ".sdf",
        )
    )


def _suffix(name: str) -> str:
    lower = name.casefold()
    for value in (".tar.gz", ".tar.bz2"):
        if lower.endswith(value):
            return value
    return Path(lower).suffix
